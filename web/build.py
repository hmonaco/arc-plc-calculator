"""Build the ARC-CO / PLC calculator page (docs/index.html).

Usage:
    python web/build.py                       # newest data/FSAInputData_*.xlsx and data/Detailed_Results_*.csv
    python web/build.py --fsa data/FSAInputData_20261012.xlsx --results data/Detailed_Results_2026_ncsa_09142026.csv

Inputs (see AUDIT.md for every field and where it is used):
    FSA workbook, tab "Main"       -> dropdowns (state / county / crop / practice), benchmark yield & price,
                                      default county yield, default MYA price, PLC yield, 2020-24 yields
    FSA workbook, tab "PLC rates"  -> effective reference price and national loan rate (matched on "name pmts")
    Detailed results CSV           -> "Our estimations" section, corn / soybeans / wheat only
                                      (matched on county FIPS + crop + practice)
    FSA yields feather             -> History chart: FSA actual county yields 2009-2025
                                      (matched on county FIPS + crop + practice)

Outputs:
    docs/index.html                -> the published page (data, font and logo embedded)
    data/app_inputs_audit.csv      -> every row the app uses, one column per field, labeled with its source
    data/build_report.txt          -> merge statistics, cross-checks and anything that did not match
    data/yield_conflicts.csv       -> yields-file records where one county/crop/practice/year has two different values
"""
import argparse
import base64
import json
import math
import re
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
WEB = Path(__file__).parent
sys.path.insert(0, str(ROOT))
from arcplc import arc_co_payment, plc_payment  # noqa: E402  (same formulas as the page, used for the audit file)

# ---------------------------------------------------------------- column map
# Main tab: header text (row 4) -> (expected column letter, internal name).
# Columns are found by HEADER TEXT. If a column moves, the build still finds it and reports the new letter;
# if a header is renamed or missing, the build stops with an error instead of reading the wrong column.
MAIN_HEADER_ROW = 4
MAIN_COLS = {
    "ST_Cty": ("A", "fips"),
    "State Name": ("B", "state"),
    "County Name": ("C", "county"),
    "Sub County": ("D", "subcty"),
    "Crop Name": ("E", "crop"),
    "name pmts": ("F", "crop_key"),
    "Unit": ("I", "unit"),
    "ARC-CO Yield Designation": ("J", "practice"),
    "{Y0} trend adjusted (county yield or 80% of T)": ("K", "y0"),
    "{Y1} trend adjusted (county yield or 80% of T)": ("L", "y1"),
    "{Y2} trend adjusted (county yield or 80% of T)": ("M", "y2"),
    "{Y3} trend adjusted (county yield or 80% of T)": ("N", "y3"),
    "{Y4} trend adjusted (county yield or 80% of T)": ("O", "y4"),
    "{PY} Bench Mark ({B0}-{B1} olympic avg)": ("Q", "bench_yield"),
    "{PY} Bench Mark Price ({B0}-{B1} olympic avg)": ("R", "bench_price"),
    "{PY} Benchmark Revenue": ("S", "bench_rev_fsa"),        # cross-check only
    "{PY} Guarantee Revenue": ("T", "guarantee_fsa"),        # cross-check only
    "{PY} Maximum Payment Rate": ("U", "max_rate_fsa"),      # cross-check only
    "{PY} Actual Yield": ("V", "default_yield"),
    "{PY} National Price": ("W", "default_price"),
    "PLC Yields": ("AA", "plc_yield"),
}
# PLC rates tab: header row is the row whose first cell is "Commodity".
RATE_COLS = {
    "Commodity": ("A", "rate_name"),
    "key": ("B", "crop_key"),
    "Unit": ("E", "rate_unit"),
    "{PY} Effective Reference Price": ("F", "erp"),
    "Projected (P) or Final (F) {PY}/{PY1} MYA Price": ("G", "rate_mya"),  # cross-check vs Main col W
    "{PY} National Loan Rate": ("I", "loan"),
}
# Header texts above use placeholders filled from the program year in the Main tab title ("... FOR PROGRAM YEAR 2026"):
# {PY}=2026, {PY1}=27, {B0}=2020, {B1}=24, {Y0}..{Y4}=2020..2024 (the five benchmark years, PY-6 .. PY-2).
# History chart years = every year present in the yields file before the program year.
PRACTICES = ["All", "Nonirrigated", "Irrigated"]
UNITS = {"Bushel": ("bu", 2), "Pound": ("lb", 4)}
# Detailed results crop names -> FSA Main "Crop Name". Only these crops get the estimations section.
EST_CROPS = {"Corn": "Corn", "Soybean": "Soybeans", "Wheat": "Wheat"}


def newest(pattern):
    files = sorted((ROOT / "data").glob(pattern))
    if not files:
        sys.exit(f"ERROR: no file matching data/{pattern}")
    return files[-1]


def program_year(path):
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)["Main"]
    title = " ".join(str(v) for r in ws.iter_rows(max_row=3, values_only=True) for v in r if v)
    mt = re.search(r"PROGRAM YEAR\s+(20\d\d)", title, re.I)
    if not mt:
        sys.exit("ERROR: could not find 'PROGRAM YEAR 20xx' in the Main tab title (rows 1-3)")
    return int(mt.group(1))


def fill_years(cols, py):
    v = {"PY": py, "PY1": str(py + 1)[2:], "B0": py - 6, "B1": str(py - 2)[2:], **{f"Y{i}": py - 6 + i for i in range(5)}}
    return {k.format(**v): val for k, val in cols.items()}


def read_sheet(path, sheet, header_row, wanted, report):
    """Read a sheet, locating each wanted header; returns DataFrame with internal names + letter map."""
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)[sheet]
    rows = list(ws.iter_rows(values_only=True))
    if header_row is None:  # find the row whose first cell is "Commodity"
        header_row = next(i + 1 for i, r in enumerate(rows) if r and str(r[0]).strip() == "Commodity")
    header = [("" if v is None else str(v).strip()) for v in rows[header_row - 1]]
    letters, idx = {}, {}
    for text, (exp_letter, name) in wanted.items():
        if text not in header:
            sys.exit(f"ERROR: column '{text}' not found in {sheet} (row {header_row}). Was it renamed?")
        i = header.index(text)
        letter = openpyxl.utils.get_column_letter(i + 1)
        letters[name] = letter
        idx[name] = i
        if letter != exp_letter:
            report.append(f"NOTE: {sheet}: '{text}' is now in column {letter} (expected {exp_letter}). Used it anyway.")
    data = [{n: r[i] for n, i in idx.items()} for r in rows[header_row:] if r and r[0] not in (None, "")]
    return pd.DataFrame(data), letters, header_row


def clean_name(s):
    return re.sub(r"\s*\d/\s*$", "", str(s)).strip()  # drop footnote markers like " 2/"


def num(x, d):
    if x is None or (isinstance(x, float) and math.isnan(x)) or x == "":
        return None
    return round(float(x), d)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--fsa", type=Path)
    ap.add_argument("--results", type=Path)
    ap.add_argument("--yields", type=Path)
    ap.add_argument("--fragment", action="store_true", help="also write docs/fragment.html (body only)")
    a = ap.parse_args()
    fsa = a.fsa or newest("FSAInputData_*.xlsx")
    res = a.results or newest("Detailed_Results_*.csv")
    yld = a.yields or newest("FSA_yields_*.feather")
    report = [f"Build {datetime.now():%Y-%m-%d %H:%M}", f"FSA workbook : {fsa.name}", f"Results file : {res.name}",
              f"Yields file  : {yld.name}", ""]

    # ---------------- FSA Main
    PY = program_year(fsa)
    HIST_YEARS = list(range(PY - 6, PY - 1))     # Main K-O: trend-adjusted yields behind the benchmark
    report.insert(4, f"Program year : {PY} (from the Main tab title); benchmark years {HIST_YEARS[0]}-{HIST_YEARS[-1]}")
    m, mL, _ = read_sheet(fsa, "Main", MAIN_HEADER_ROW, fill_years(MAIN_COLS, PY), report)
    m["fips"] = m["fips"].astype(str).str.zfill(5)
    m["subcty"] = m["subcty"].fillna("").astype(str).str.strip()
    for c in ["state", "county", "crop", "crop_key", "unit", "practice"]:
        m[c] = m[c].astype(str).str.strip()
    numeric = [n for n in m.columns if n not in ("fips", "state", "county", "subcty", "crop", "crop_key", "unit", "practice")]
    for c in numeric:
        m[c] = pd.to_numeric(m[c], errors="coerce")
    report.append(f"Main tab rows read: {len(m)}  (crops: {m.crop.nunique()}, counties: {m[['fips','subcty']].drop_duplicates().shape[0]})")
    dup = m.duplicated(["fips", "subcty", "crop", "practice"]).sum()
    if dup:
        sys.exit(f"ERROR: {dup} duplicate county/sub-county/crop/practice rows in Main")
    bad_prac = set(m.practice) - set(PRACTICES)
    if bad_prac:
        sys.exit(f"ERROR: unexpected practice values in Main col {mL['practice']}: {bad_prac}")
    bad_unit = set(m.unit) - set(UNITS)
    if bad_unit:
        sys.exit(f"ERROR: unexpected units in Main col {mL['unit']}: {bad_unit}")

    # ---------------- FSA PLC rates
    rt, rL, rrow = read_sheet(fsa, "PLC rates", None, fill_years(RATE_COLS, PY), report)
    rt = rt[rt["crop_key"].notna() & pd.to_numeric(rt["erp"], errors="coerce").notna()].copy()
    rt["crop_key"] = rt["crop_key"].astype(str).str.strip()
    for c in ["erp", "rate_mya", "loan"]:
        rt[c] = pd.to_numeric(rt[c])
    ws = openpyxl.load_workbook(fsa, read_only=True, data_only=True)["PLC rates"]
    rates_date = next((str(r[0]).replace("1/", "").strip() for r in ws.iter_rows(max_row=5, values_only=True)
                       if r[0] and re.search(r"20\d\d", str(r[0])) and len(str(r[0])) < 40), "")
    report.append(f"PLC rates tab: {len(rt)} commodities (header row {rrow}), dated '{rates_date}'")
    missing_keys = set(m.crop_key) - set(rt.crop_key)
    if missing_keys:
        sys.exit(f"ERROR: Main 'name pmts' values with no PLC rates row: {missing_keys}")
    m = m.merge(rt[["crop_key", "rate_name", "erp", "loan", "rate_mya"]], on="crop_key", how="left")

    # ---------------- cross-checks (reported, not fatal)
    report.append("")
    report.append("CROSS-CHECKS (FSA file internal consistency)")
    br = m.bench_yield * m.bench_price
    chk = [
        ("Benchmark revenue S = Q x R", (br - m.bench_rev_fsa).abs(), 0.01),
        ("Guarantee T = 90% of S", (0.9 * m.bench_rev_fsa - m.guarantee_fsa).abs(), 0.01),
        ("Max payment U = 12% of S", (0.12 * m.bench_rev_fsa - m.max_rate_fsa).abs(), 0.01),
        ("National price W = PLC rates projected MYA (G)", (m.default_price - m.rate_mya).abs(), 1e-9),
    ]
    for label, diff, tol in chk:
        n_bad = int((diff > tol).sum())
        report.append(f"  {label:<48} max diff {diff.max():.4f}  rows over tolerance: {n_bad}")
    no_bp = m[m.bench_price.isna()]
    if len(no_bp):
        report.append(f"  Rows with no benchmark price (ARC-CO shown as n/a): {len(no_bp)} -> {sorted(no_bp.crop.unique())}")

    # ---------------- detailed results (estimations) for corn / soybeans / wheat
    d = pd.read_csv(res)
    d = d[d.Commodity.isin(EST_CROPS) & d.YieldType.isin(PRACTICES)].copy()
    d["crop"] = d.Commodity.map(EST_CROPS)
    d["fips"] = d.fips.astype(int).astype(str).str.zfill(5)
    d = d.rename(columns={"YieldType": "practice", "Avg. Benchmark Yield": "est_bench_yield",
                          f"{PY} Avg. ARC-CO Pmt": "est_arc", f"{PY} Avg. PLC Pmt": "est_plc",
                          f"{PY} ARC-CO Pmt Likelihood": "est_p_arc", f"{PY} PLC Pmt Likelihood": "est_p_plc"})
    keep = ["fips", "crop", "practice", "est_bench_yield", "est_arc", "est_plc", "est_p_arc", "est_p_plc"]
    m["_est_eligible"] = m.crop.isin(EST_CROPS.values()) & (m["subcty"] == "")
    m = m.merge(d[keep], on=["fips", "crop", "practice"], how="left")
    m.loc[~m._est_eligible, keep[3:]] = float("nan")
    elig = m[m._est_eligible]
    matched = elig.est_arc.notna()
    report.append("")
    report.append("ESTIMATIONS MERGE (results file -> FSA Main; key = county FIPS + crop + practice)")
    report.append(f"  FSA corn/soybean/wheat rows (no sub-county): {len(elig)}; matched: {int(matched.sum())}; unmatched: {int((~matched).sum())}")
    sub_est = m[m.crop.isin(EST_CROPS.values()) & (m['subcty'] != '')]
    if len(sub_est):
        report.append(f"  Sub-county rows skipped (results file has no sub-county): {len(sub_est)} "
                      f"({', '.join(sorted(set(sub_est.county + ' ' + sub_est['subcty'])))})")
    d_keys = set(map(tuple, d[["fips", "crop", "practice"]].values))
    m_keys = set(map(tuple, m[["fips", "crop", "practice"]].values))
    extra = sorted(d_keys - m_keys)
    report.append(f"  Results rows with no FSA row (not shown in app): {len(extra)}")
    for k in extra:
        report.append(f"      {k[0]} {k[1]} {k[2]}")
    for k in elig[~matched][["fips", "county", "crop", "practice"]].values:
        report.append(f"  FSA row with no results row (section hidden): {' '.join(map(str, k))}")
    m["est_bench_diff"] = (m.est_bench_yield - m.bench_yield).round(2)
    mism = m[m.est_bench_diff.abs() > 0.05]
    report.append(f"  Benchmark yield differs between results file and FSA col {mL['bench_yield']} (> 0.05): {len(mism)}")
    for r in mism.itertuples():
        report.append(f"      {r.fips} {r.county}, {r.state} | {r.crop} {r.practice}: FSA {r.bench_yield}  results {r.est_bench_yield}")

    # ---------------- FSA actual yields (History chart) -- key = county FIPS + crop + practice
    y = pd.read_feather(yld).reset_index()
    need = ["ST_Cty", "Crop Name", "ARC-CO Yield Designation", "Year", "ActYield"]
    miss = [c for c in need if c not in y.columns]
    if miss:
        sys.exit(f"ERROR: yields file is missing columns {miss}")
    y = y[need].rename(columns={"ST_Cty": "fips", "Crop Name": "crop", "ARC-CO Yield Designation": "practice",
                                "Year": "year", "ActYield": "act"})
    y["fips"] = y.fips.astype(int).astype(str).str.zfill(5)
    ky = ["fips", "crop", "practice", "year"]
    grp = y.groupby(ky).act.agg(["count", "nunique"])
    conf = grp[grp["nunique"] > 1].index
    report.append("")
    report.append("YIELDS FILE (History chart; key = county FIPS + crop + practice + year)")
    report.append(f"  Records: {len(y)}; duplicated keys: {int((grp['count'] > 1).sum())}; with conflicting values: {len(conf)}")
    if len(conf):
        cf = y.set_index(ky).loc[conf].reset_index()
        cf["order_in_file"] = cf.groupby(ky).cumcount() + 1
        cf.to_csv(ROOT / "data" / "yield_conflicts.csv", index=False)
        report.append("  Rule: the FIRST value in file order is used. All conflicting values are listed in data/yield_conflicts.csv")
        report.append("  Conflicts by year: " + ", ".join(f"{k}: {v}" for k, v in cf.drop_duplicates(ky).year.value_counts().sort_index().items()))
    y = y.drop_duplicates(ky, keep="first")
    ACT_YEARS = sorted(int(v) for v in y.year.unique() if v < PY)
    report.append(f"  Years shown in the History chart: {ACT_YEARS[0]}-{ACT_YEARS[-1]}")
    act = y[y.year.isin(ACT_YEARS)].pivot_table(index=["fips", "crop", "practice"], columns="year", values="act", aggfunc="first")
    act = act.reindex(columns=ACT_YEARS)
    act.columns = [f"act{c}" for c in act.columns]
    m = m.merge(act.reset_index(), on=["fips", "crop", "practice"], how="left")
    m.loc[m["subcty"] != "", act.columns] = float("nan")   # yields file has no sub-county
    no_hist = m[m[act.columns].isna().all(axis=1)]
    report.append(f"  Main rows with yield history: {len(m) - len(no_hist)} of {len(m)}; without: {len(no_hist)} "
                  f"(of which sub-county rows: {int((no_hist.subcty != '').sum())})")
    for r in no_hist[no_hist.subcty == ""].itertuples():
        report.append(f"      no history: {r.fips} {r.county}, {r.state} | {r.crop} {r.practice}")

    # ---------------- page payload
    lk = pd.read_excel(fsa, sheet_name="lookups")
    pretty = {str(r.iloc[2]).strip(): clean_name(r.iloc[0]) for _, r in lk.iterrows() if pd.notna(r.iloc[2])}
    crops = sorted(m.crop.unique(), key=lambda c: pretty.get(c, c))
    cidx = {c: i for i, c in enumerate(crops)}
    crop_info = []
    for c in crops:
        g = m[m.crop == c].iloc[0]
        u, dec = UNITS[g.unit]
        crop_info.append({"name": pretty.get(c, c), "fsa": c, "unit": u, "dec": dec,
                          "erp": num(g.erp, 6), "loan": num(g.loan, 6)})

    places = m[["fips", "subcty", "state", "county"]].drop_duplicates().sort_values(["state", "county", "subcty"]).reset_index(drop=True)
    pidx = {(f, sc): i for i, (f, sc) in enumerate(zip(places.fips, places.subcty))}
    postal = pd.read_csv(ROOT / "data" / "counties.csv").assign(sf=lambda x: x.fips // 1000).drop_duplicates("sf").set_index("sf").postal.to_dict()

    rows = []
    for r in m.itertuples():
        est = None
        if pd.notna(r.est_arc):
            est = [num(r.est_arc, 2), num(r.est_plc, 2), num(r.est_p_arc, 3), num(r.est_p_plc, 3), num(r.est_bench_yield, 2)]
        rows.append([pidx[(r.fips, r.subcty)], cidx[r.crop], PRACTICES.index(r.practice),
                     [num(getattr(r, f"act{yr}"), 2) for yr in ACT_YEARS],
                     num(r.bench_yield, 6), num(r.bench_price, 6), num(r.plc_yield, 6),
                     num(r.default_yield, 6), num(r.default_price, 6), est])  # full precision for anything used in a formula
    data = {
        "year": PY, "benchYears": [HIST_YEARS[0], HIST_YEARS[-1]], "histYears": ACT_YEARS, "practices": PRACTICES, "crops": crop_info,
        "places": [[int(p.fips), p.subcty, p.state, p.county, postal.get(int(p.fips) // 1000, "")] for p in places.itertuples()],
        "rows": rows,
        "sources": {"fsa": fsa.name, "results": res.name, "rates_date": rates_date},
    }

    # ---------------- audit table (exactly what the page uses, with sources)
    a_df = pd.DataFrame({
        "County FIPS [Main!A]": m.fips, "State [Main!B]": m.state, "County [Main!C]": m.county,
        "Sub county [Main!D]": m["subcty"], "Crop [Main!E]": m.crop, "PLC rates key [Main!F]": m.crop_key,
        "Unit [Main!I]": m.unit, "Practice [Main!J]": m.practice,
        **{f"{y} trend-adj. yield [Main!{mL['y'+str(i)]}]": m[f"y{i}"] for i, y in enumerate(HIST_YEARS)},
        **{f"{yr} actual yield [yields file]": m[f"act{yr}"] for yr in ACT_YEARS},
        f"Benchmark yield [Main!{mL['bench_yield']}]": m.bench_yield,
        f"Benchmark price [Main!{mL['bench_price']}]": m.bench_price,
        f"Default county yield [Main!{mL['default_yield']}]": m.default_yield,
        f"Default MYA price [Main!{mL['default_price']}]": m.default_price,
        f"PLC yield [Main!{mL['plc_yield']}]": m.plc_yield,
        f"Eff. reference price [PLC rates!{rL['erp']}]": m.erp,
        f"Loan rate [PLC rates!{rL['loan']}]": m.loan,
        "ARC-CO at defaults $/base ac [computed]": arc_co_payment(m.default_yield, m.default_price, m.bench_yield, m.bench_price, m.loan).round(2),
        "PLC at defaults $/base ac [computed]": plc_payment(m.default_price, m.erp, m.loan, m.plc_yield).round(2),
        "Est. ARC-CO [results]": m.est_arc, "Est. PLC [results]": m.est_plc,
        "Chance ARC-CO pays [results]": m.est_p_arc, "Chance PLC pays [results]": m.est_p_plc,
        "Results benchmark yield [results]": m.est_bench_yield,
        "Benchmark yield diff results-FSA": m.est_bench_diff,
    })
    a_df.to_csv(ROOT / "data" / "app_inputs_audit.csv", index=False)

    # ---------------- assemble page
    blob = json.dumps(data, separators=(",", ":"))
    html = (WEB / "template.html").read_text(encoding="utf-8").replace("/*__DATA__*/null", blob)
    b64 = lambda f: base64.b64encode(f.read_bytes()).decode()
    if (WEB / "logo.png").exists():
        imgs = f'<img class="logo-light" src="data:image/png;base64,{b64(WEB / "logo.png")}" alt="University of Minnesota Extension">'
        if (WEB / "logo-dark.png").exists():
            imgs += f'<img class="logo-dark" src="data:image/png;base64,{b64(WEB / "logo-dark.png")}" alt="University of Minnesota Extension">'
        html = html.replace('<!--__LOGO__--><div class="wordmark">', imgs + '<div class="wordmark" hidden>', 1)
    faces = []
    for f in sorted((WEB / "fonts").glob("*-latin-*-normal.woff2")):
        fam, _, wt = f.stem.split("-")[:3]
        faces.append(f'@font-face{{font-family:"{fam.title()}";font-style:normal;font-weight:{wt};font-display:swap;'
                     f'src:url(data:font/woff2;base64,{b64(f)}) format("woff2")}}')
    html = html.replace("/*__FONTS__*/", "\n".join(faces), 1)
    page = ('<!doctype html>\n<html lang="en">\n<head>\n<meta charset="utf-8">\n'
            '<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">\n'
            '<style>body{margin:0}[hidden]{display:none!important}</style>\n</head>\n<body>\n'
            + html + "\n</body>\n</html>\n")
    out = ROOT / "docs" / "index.html"
    out.parent.mkdir(exist_ok=True)
    out.write_text(page, encoding="utf-8")
    if a.fragment:
        out.with_name("fragment.html").write_text(html, encoding="utf-8")

    report.append("")
    report.append(f"Wrote docs/index.html ({out.stat().st_size / 1e6:.2f} MB, {len(rows)} rows, {len(places)} counties/sub-counties)")
    report.append("Wrote data/app_inputs_audit.csv")
    (ROOT / "data" / "build_report.txt").write_text("\n".join(report) + "\n", encoding="utf-8")
    print("\n".join(report[:3] + [l for l in report if l.startswith(("NOTE", "  FSA corn", "  Benchmark", "Wrote"))]))


if __name__ == "__main__":
    main()
